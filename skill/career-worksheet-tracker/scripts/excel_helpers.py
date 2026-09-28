# -*- coding: utf-8 -*-
"""
진로활동지 스캔본 -> 학생별 누적 엑셀 정리 스킬의 재사용 헬퍼.

사용 예:
    from excel_helpers import (
        new_workbook, get_or_create_student_sheet, append_row, save_and_reorder
    )

    wb = new_workbook()  # 또는 openpyxl.load_workbook(path) 로 기존 파일 열기
    ws = get_or_create_student_sheet(wb, "10101", "정솔미", grade_label="2026학년도 1학년")
    append_row(ws, {
        "차시": "1-1",
        "활동명": "진로수업 OT",
        "개요": "자기소개, 진로활동 안내, 진로수업 약속 정하기",
        "응답요약": "...",
        "특이사항": "...",
        "비고": "...",
    })
    save_and_reorder(wb, "/mnt/user-data/outputs/1학년_1반_진로활동정리.xlsx")

진로 프로젝트 진행 확인표:
    data = read_progress(wb)                      # 기존 표 읽기(없으면 {})
    update_progress(data, "10101", name="정솔미", sess_no=3, submitted="○", level="상",
                    note_add="P-3 자료 4개 충실히 요약", track="주제탐구", topic="...")
    write_progress_sheet(wb, data, "2026학년도 1학년 1반 진로 프로젝트 진행 확인표")
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = "맑은 고딕"

HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
BASE_FONT = Font(name=FONT_NAME, size=10)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True)
SUB_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

COLUMN_ORDER = ["차시", "활동명", "개요", "응답요약", "특이사항", "비고"]
HEADERS = ["차시", "활동명", "학습목표/개요", "학생 응답 요약", "특이사항", "비고(미기재/확인필요)"]
COLUMN_WIDTHS = {"A": 8, "B": 26, "C": 34, "D": 55, "E": 26, "F": 30}
DEFAULT_ROW_HEIGHT = 140

DEFAULT_NOTE = (
    "※ 활동지 스캔본(PDF)을 바탕으로 정리. 필체 판독이 어렵거나 학생이 미기재한 항목은 "
    "'판독 어려움/미기재'로 표기함."
)

SUMMARY_SHEET_NAME = "차시 목록"


def new_workbook():
    """빈 워크북을 만들고 기본 시트를 지운다."""
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)
    return wb


def sheet_name_for(student_id: str, student_name: str) -> str:
    return f"{student_id} {student_name}"


def find_student_sheet(wb, student_id: str):
    """학번 접두어로 기존 시트를 찾는다. 없으면 None."""
    for name in wb.sheetnames:
        if name.startswith(student_id) and " " in name:
            return wb[name]
    return None


def get_or_create_student_sheet(
    wb, student_id: str, student_name: str, grade_label: str = "", note: str = None
):
    """학생 시트가 있으면 반환, 없으면 제목/안내/헤더까지 세팅한 새 시트를 만든다."""
    existing = find_student_sheet(wb, student_id)
    if existing is not None:
        return existing

    name = sheet_name_for(student_id, student_name)
    ws = wb.create_sheet(name)

    title = f"{grade_label} 진로수업 차시별 활동 정리 (학번 {student_id} {student_name})".strip()
    ws.merge_cells("A1:F1")
    c = ws.cell(row=1, column=1, value=title)
    c.font = TITLE_FONT

    ws.merge_cells("A2:F2")
    c = ws.cell(row=2, column=1, value=note or DEFAULT_NOTE)
    c.font = SUB_FONT

    header_row = 4
    for i, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A5"

    return ws


def append_row(ws, row_dict: dict, row_height: int = DEFAULT_ROW_HEIGHT):
    """
    기존 표 맨 아래(ws.max_row + 1)에 헤더와 같은 스타일로 한 행을 추가한다.
    학기가 바뀌어도 별도 섹션 헤더를 넣지 말고 이 함수만 호출해서 이어붙일 것.
    row_dict 키: 차시, 활동명, 개요, 응답요약, 특이사항, 비고
    """
    row_idx = ws.max_row + 1
    values = [row_dict.get(k, "") for k in COLUMN_ORDER]
    for col_idx, v in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=v)
        cell.font = BASE_FONT
        cell.alignment = WRAP
        cell.border = BORDER
    ws.cell(row=row_idx, column=1).alignment = WRAP_CENTER
    ws.row_dimensions[row_idx].height = row_height
    return row_idx


def ensure_summary_sheet_last(wb, summary_sheet_name: str = SUMMARY_SHEET_NAME):
    """이름이 summary_sheet_name으로 시작하는 시트를 찾아 맨 끝으로 옮긴다."""
    for name in wb.sheetnames:
        if name.startswith(summary_sheet_name):
            wb.move_sheet(name, offset=len(wb.sheetnames))
            break


def save_and_reorder(wb, path: str, summary_sheet_name: str = SUMMARY_SHEET_NAME):
    ensure_summary_sheet_last(wb, summary_sheet_name)
    wb.save(path)
    return path


def find_all_student_ids(wb):
    """워크북 안의 모든 학생 학번을 시트명에서 추출 (요약 시트 제외)."""
    ids = []
    for name in wb.sheetnames:
        if " " in name and name.split(" ")[0].isdigit():
            ids.append(name.split(" ")[0])
    return ids


# ---------------------------------------------------------------------------
# 진로 프로젝트 진행 확인표 (학생 1명 = 1줄, 차시별 제출 ○/× + 기록정도 상/중/하)
# ---------------------------------------------------------------------------
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as _L

PROGRESS_SHEET_NAME = "프로젝트 진행 확인표"
PROGRESS_N_SESS = 8
PROGRESS_FIRST_COL = 5   # P-1 '제출' 열(E)
PROGRESS_FIRST_ROW = 6   # 첫 학생 행
PROGRESS_FILLS = {"상": "C6EFCE", "중": "FFEB9C", "하": "F8CBAD", "×": "D9D9D9"}
PROGRESS_NOTE = (
    "※ 제출: ○ 제출 / × 미제출(빈칸 제출도 ○로 표시하고 기록정도 '하')   "
    "※ 기록정도 — 상: 모든 항목을 구체적으로 작성(주제탐구는 자료 3개 이상+요약) / "
    "중: 핵심 항목은 작성했으나 일부 빈칸·자료 부족 / 하: 대부분 빈칸이거나 출처·한 줄만 작성   "
    "※ 이후 차시 칸은 드롭다운으로 직접 입력하면 색이 자동으로 칠해짐"
)


def read_progress(wb, n_sess: int = PROGRESS_N_SESS):
    """
    기존 진행 확인표를 dict로 읽는다. 시트가 없으면 빈 dict.
    반환: {학번: {"name", "track", "topic", "sess": {1: ("○","상"), ...}, "note"}}
    """
    if PROGRESS_SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[PROGRESS_SHEET_NAME]
    note_col = PROGRESS_FIRST_COL + 2 * n_sess + 2
    data = {}
    r = PROGRESS_FIRST_ROW
    while True:
        sid = ws.cell(r, 1).value
        if sid is None or not str(sid).isdigit():
            break
        sess = {}
        for i in range(n_sess):
            sub = ws.cell(r, PROGRESS_FIRST_COL + 2 * i).value
            lv = ws.cell(r, PROGRESS_FIRST_COL + 2 * i + 1).value
            if sub or lv:
                sess[i + 1] = (sub, lv)
        data[str(sid)] = {
            "name": ws.cell(r, 2).value, "track": ws.cell(r, 3).value,
            "topic": ws.cell(r, 4).value, "sess": sess, "note": ws.cell(r, note_col).value or "",
        }
        r += 1
    return data


def update_progress(data, sid, name=None, sess_no=None, submitted=None, level=None,
                    note_add=None, track=None, topic=None):
    """
    data(read_progress 결과)에 한 학생·한 차시 기록을 반영한다.
    note_add는 간단기록 끝에 ' → '로 이어 붙인다(예: 'P-3 자료 4개 충실히 요약').
    """
    rec = data.setdefault(str(sid), {"name": name or "", "track": "", "topic": "", "sess": {}, "note": ""})
    if name: rec["name"] = name
    if track: rec["track"] = track
    if topic: rec["topic"] = topic
    if sess_no:
        rec["sess"][sess_no] = (submitted, None if submitted == "×" else level)
    if note_add:
        rec["note"] = f"{rec['note']} → {note_add}" if rec["note"] else note_add
    return rec


def write_progress_sheet(wb, data, title: str, n_sess: int = PROGRESS_N_SESS,
                         summary_sheet_name: str = SUMMARY_SHEET_NAME):
    """
    진행 확인표 시트를 data로 새로 그린다(기존 시트는 지우고 다시 만듦 → 새 학생 추가·정렬도 자동).
    title 예: '2026학년도 1학년 12반 진로 프로젝트 진행 확인표'
    시트 순서는 학생 시트들 → 진행 확인표 → 요약 시트(맨 끝).
    """
    from openpyxl.styles import Font as _F, Alignment as _A, PatternFill as _P
    if PROGRESS_SHEET_NAME in wb.sheetnames:
        del wb[PROGRESS_SHEET_NAME]
    ws = wb.create_sheet(PROGRESS_SHEET_NAME)
    fc = PROGRESS_FIRST_COL
    last_col = fc + 2 * n_sess - 1
    cnt_col, hi_col, note_col = last_col + 1, last_col + 2, last_col + 3
    center = _A(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=note_col)
    ws.cell(1, 1, title).font = TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=note_col)
    c = ws.cell(2, 1, PROGRESS_NOTE)
    c.font = _F(name=FONT_NAME, size=9, italic=True, color="595959"); c.alignment = WRAP
    ws.row_dimensions[2].height = 42

    for col, h in {1: "학번", 2: "이름", 3: "트랙", 4: "주제 / 도서 / 습관"}.items():
        ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        ws.cell(4, col, h)
    for i in range(n_sess):
        col = fc + 2 * i
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.cell(4, col, f"P-{i+1}")
        ws.cell(5, col, "제출"); ws.cell(5, col + 1, "기록")
    for col, h in [(cnt_col, "제출\n차시 수"), (hi_col, "'상'\n차시 수"),
                   (note_col, "간단기록(차시별 흐름 · 지도 포인트)")]:
        ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        ws.cell(4, col, h)
    for r in (4, 5):
        for col in range(1, note_col + 1):
            cell = ws.cell(r, col)
            cell.font = HEADER_FONT; cell.fill = HEADER_FILL; cell.alignment = center; cell.border = BORDER

    r0 = PROGRESS_FIRST_ROW
    sids = sorted(data)
    for i, sid in enumerate(sids):
        r = r0 + i
        rec = data[sid]
        vals = {1: int(sid), 2: rec.get("name"), 3: rec.get("track"), 4: rec.get("topic"),
                note_col: rec.get("note")}
        for k, (sub, lv) in rec.get("sess", {}).items():
            vals[fc + 2 * (k - 1)] = sub
            vals[fc + 2 * (k - 1) + 1] = lv or None
        for col in range(1, note_col + 1):
            cell = ws.cell(r, col, vals.get(col))
            cell.font = BASE_FONT; cell.border = BORDER
            cell.alignment = WRAP if col in (4, note_col) else center
        ws.cell(r, cnt_col, f'=COUNTIF({_L(fc)}{r}:{_L(last_col)}{r},"○")')
        ws.cell(r, hi_col, f'=COUNTIF({_L(fc)}{r}:{_L(last_col)}{r},"상")')
        ws.row_dimensions[r].height = 36
    r_last = r0 + len(sids) - 1

    total_fill = _P("solid", start_color="D9E1F2")
    for j, (lab, key, is_sub) in enumerate([("제출 인원(○)", "○", True), ("미제출(×)", "×", True),
                                            ("상", "상", False), ("중", "중", False), ("하", "하", False)]):
        r = r_last + 2 + j
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(r, 1, lab)
        for col in range(1, note_col + 1):
            cell = ws.cell(r, col)
            cell.border = BORDER; cell.fill = total_fill; cell.alignment = center
            cell.font = _F(name=FONT_NAME, size=10, bold=True)
        for i in range(n_sess):
            col = fc + 2 * i + (0 if is_sub else 1)
            rng = f"{_L(col)}{r0}:{_L(col)}{r_last}"
            ws.cell(r, col, f'=IF(COUNTA({rng})=0,"",COUNTIF({rng},"{key}"))')

    area = f"{_L(fc)}{r0}:{_L(last_col)}{r_last}"
    for v, color in PROGRESS_FILLS.items():
        ws.conditional_formatting.add(area, CellIsRule(
            operator="equal", formula=[f'"{v}"'], fill=_P("solid", start_color=color, end_color=color)))
    dv_s = DataValidation(type="list", formula1='"○,×"', allow_blank=True)
    dv_l = DataValidation(type="list", formula1='"상,중,하"', allow_blank=True)
    ws.add_data_validation(dv_s); ws.add_data_validation(dv_l)
    for i in range(n_sess):
        dv_s.add(f"{_L(fc+2*i)}{r0}:{_L(fc+2*i)}{r_last}")
        dv_l.add(f"{_L(fc+2*i+1)}{r0}:{_L(fc+2*i+1)}{r_last}")

    for col, w in {"A": 8, "B": 12, "C": 10, "D": 30}.items():
        ws.column_dimensions[col].width = w
    for col in range(fc, last_col + 1):
        ws.column_dimensions[_L(col)].width = 5.5
    ws.column_dimensions[_L(cnt_col)].width = 8
    ws.column_dimensions[_L(hi_col)].width = 8
    ws.column_dimensions[_L(note_col)].width = 70
    ws.freeze_panes = ws.cell(r0, fc)
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = "70AD47"

    wb.move_sheet(PROGRESS_SHEET_NAME, offset=len(wb.sheetnames))
    ensure_summary_sheet_last(wb, summary_sheet_name)
    return ws
